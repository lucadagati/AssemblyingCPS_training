#!/usr/bin/env python3
"""
English slide helpers — image allowlist and image slide builder.
Deck content lives in training/decks/*.py
Regenerate all: .venv/bin/python generate_slides_modular_en.py
"""
from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from generate_slides import (
    ASSETS,
    BASE,
    GRAY,
    NAVY,
    OFF_WHITE,
    TEAL,
    TEAL_LIGHT,
    WHITE,
    _footer,
    _set_font,
    _title_card,
    _top_accent_bar,
    add_content_slide,
    blank_slide,
    new_prs,
)

VM_IP = "{{VM_IP}}"

# Max image area (inches) — keeps screenshots/diagrams inside slide + caption + footer
IMG_MAX_W = 10.8
IMG_MAX_H = 4.35
IMG_FRAME_PAD = 0.08

ALLOWED_IMAGES = {
    "chapter13/horizon-login.png",
    "chapter13/horizon-after-login.png",
    "chapter13/horizon-boards-dashboard.png",
    "chapter13/lr-ui-login.png",
    "chapter13/lr-dashboard-home.png",
    "chapter13/lr-dashboard-status.png",
    "chapter13/lr-dashboard-conf.png",
    "chapter13/lr-dashboard-system.png",
    "chapter13/sshot-create-board.png",
    "chapter13/sshot-lr-conf.png",
    "chapter13/sshot-board-info.png",
    "chapter14/horizon-plugins-dashboard.png",
    "chapter14/horizon-fleets-dashboard.png",
    "chapter14/horizon-webservices-dashboard.png",
    "chapter14/horizon-boards-with-services.png",
    "chapter14/horizon-services-list.png",
    "chapter14/horizon-service-create-filled.png",
    "chapter14/weather-server-dashboard.png",
    "chapter14/weather-sensors-json.png",
    "chapter14/wstun-nginx-success.png",
    "diagrams/wstun-port-forwarding.png",
    "diagrams/s4t-stack.png",
    "diagrams/plugin-sync-async.png",
    "diagrams/environmental-dataflow.png",
    "chapter15/influxdb-debug.png",
    "chapter15/influx-query-environmental.png",
    "chapter19/lr-dashboard-home-lr2.png",
    "chapter19/lr-dashboard-conf-lr2.png",
    "chapter19/lr-dashboard-home-lr3.png",
    "chapter19/lr-dashboard-conf-lr3.png",
    "diagrams/multiboard-fleet.png",
    "diagrams/fl-architecture.png",
    "diagrams/blueprint-layers.png",
    "diagrams/faas-contrast.png",
    "diagrams/vn-attach-workflow.png",
    "diagrams/seq-training-workflow.png",
    "diagrams/seq-board-onboarding.png",
    "diagrams/seq-plugin-sync.png",
    "diagrams/seq-environmental-async.png",
    "diagrams/seq-wstun-enable.png",
    "diagrams/seq-fleet-inject.png",
    "diagrams/seq-fl-round.png",
}


def _fit_image_inches(path, max_w: float, max_h: float) -> tuple[float, float]:
    """Return width/height in inches preserving aspect ratio."""
    from PIL import Image

    with Image.open(path) as im:
        px_w, px_h = im.size
    if px_w <= 0 or px_h <= 0:
        return max_w, max_h
    aspect = px_w / px_h
    w_in = max_w
    h_in = w_in / aspect
    if h_in > max_h:
        h_in = max_h
        w_in = h_in * aspect
    return w_in, h_in


def add_image_slide(prs, title, image_rel, caption="", notes="", hands_on=False):
    if image_rel not in ALLOWED_IMAGES or not (ASSETS / image_rel).is_file():
        return add_content_slide(
            prs,
            title,
            ["Screenshot not available", caption] if caption else ["Screenshot not available"],
            notes=notes,
            hands_on=hands_on,
        )

    slide = blank_slide(prs, bg=OFF_WHITE)
    bar_color = TEAL if image_rel.startswith("diagrams/") else NAVY
    _top_accent_bar(slide, bar_color)
    top = _title_card(slide, title, Inches(0.48), color=NAVY, size=22)

    img_path = ASSETS / image_rel
    w_in, h_in = _fit_image_inches(img_path, IMG_MAX_W, IMG_MAX_H)
    frame_w = w_in + IMG_FRAME_PAD * 2
    frame_h = h_in + IMG_FRAME_PAD * 2
    frame_left = (13.333 - frame_w) / 2
    img_top = top + Inches(0.08)

    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(frame_left),
        img_top,
        Inches(frame_w),
        Inches(frame_h),
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = TEAL_LIGHT
    frame.line.width = Pt(1.0)

    slide.shapes.add_picture(
        str(img_path),
        Inches(frame_left + IMG_FRAME_PAD),
        img_top + Inches(IMG_FRAME_PAD),
        width=Inches(w_in),
        height=Inches(h_in),
    )

    if caption:
        cap = slide.shapes.add_textbox(
            Inches(0.65), img_top + Inches(frame_h + 0.15), Inches(12), Inches(0.45)
        )
        cp = cap.text_frame.paragraphs[0]
        cp.text = caption
        _set_font(cp, size=11, color=GRAY)

    tag = "Diagram" if image_rel.startswith("diagrams/") else "Screenshot"
    _footer(slide, tag)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def main():
    from generate_slides_modular_en import main as build_all

    build_all()


if __name__ == "__main__":
    main()
